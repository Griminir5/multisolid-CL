"""Report selections and real Excel output, without running a solver."""

from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
import xarray as xr
from openpyxl import load_workbook

from packed_bed.report_schema import axis_description, describe_dataset, describe_inputs
from packed_bed.reports import ModelField, ReportSpec, reporting_times
from packed_bed_ui import workbook
from packed_bed_ui.project import Project, write_documents, write_json
from packed_bed_ui.workbook import (
    ExportCancelled, coordinate_index, expand_columns, information_rows,
    plan_sheet, plan_workbook, run_information, table_rows, write_workbook,
)


def column(quantity, **fixed):
    return {"quantity": quantity, "fixed": {d: {"value": v} for d, v in fixed.items()}}


def sheet(axis="time", columns=None, name="Data", rows=None):
    return {"name": name, "axis": axis, "columns": columns or [],
            "rows": rows or {"mode": "all"}}


def report(*sheets):
    return {"version": 1, "sheets": list(sheets)}


@pytest.fixture
def dataset():
    coords = {"time": ("time", [0., 1., 2.], {"units": "s"}),
              "x_cell": ("x_cell", [.15, .45, .75], {"units": "m"}),
              "x_face": ("x_face", [0., .3, .6, .9], {"units": "m"}),
              "gas_species": ["CH4", "CO", "H2"], "solid_species": ["Ni", "NiO"],
              "delta": ("delta", [0., .2], {"long_name": "Oxygen deficiency", "units": "1"})}
    return xr.Dataset({
        "gas_mole_fraction": (("time", "gas_species", "x_cell"), np.arange(27).reshape(3, 3, 3) / 100, {"units": "1"}),
        "temperature": (("time", "x_cell"), np.arange(9).reshape(3, 3) + 300., {"units": "K"}),
        "solid_mole_fraction": (("time", "solid_species", "x_cell"), np.arange(18).reshape(3, 2, 3) / 20, {"units": "1"}),
        "velocity": (("time", "x_face"), np.arange(12).reshape(3, 4) / 10, {"units": "m/s"}),
        "outlet_temperature": ("time", [301., 304., 307.], {"units": "K"}),
        "distribution": (("time", "x_cell", "gas_species", "delta"), np.arange(54).reshape(3, 3, 3, 2), {"units": "mol"}),
    }, coords=coords)


@pytest.fixture
def recorded(tmp_path, source_case, dataset):
    project = Project.create(tmp_path / "project")
    case = project.add_case_from_files(source_case, "Original case")
    write_documents(case.run_folder / "inputs", case.documents)
    write_json(case.run_folder / "snapshot.json", {
        "case_id": case.id, "case_name": case.name, "attempt_id": "attempt-one",
        "fingerprint": case.fingerprint(), "input_hashes": {"run": "retained-hash"},
    })
    write_json(case.run_folder / "status.json", {"state": "completed"})
    (case.run_folder / "output").mkdir()
    dataset.to_netcdf(case.run_folder / "output/results.nc", engine="scipy")
    return case


def test_time_profile_and_additional_axis(dataset):
    schema = describe_dataset(dataset)
    columns = expand_columns("gas_mole_fraction", {
        "gas_species": [{"value": v} for v in ["CH4", "CO", "H2"]],
        "x_cell": [{"value": v} for v in [.15, .45, .75]],
    })
    assert len(columns) == 9
    columns = [columns[8], {**columns[0], "label": "Reference"}]
    columns += [column("outlet_temperature"), column("velocity", x_face=.9)]
    table = plan_sheet(schema, sheet(columns=columns))
    assert len(table["headings"]) == 5
    rows = np.asarray(list(table_rows(dataset, table)))
    np.testing.assert_equal(rows[:, 1], dataset.gas_mole_fraction.sel(gas_species="H2", x_cell=.75))
    np.testing.assert_equal(rows[:, 2], dataset.gas_mole_fraction.sel(gas_species="CH4", x_cell=.15))
    np.testing.assert_equal(rows[:, 3], dataset.outlet_temperature)
    np.testing.assert_equal(rows[:, 4], dataset.velocity.sel(x_face=.9))
    assert table["headings"][2] == "Reference (1)"

    profile = plan_sheet(schema, sheet("x_cell", [column("temperature", time=1.),
                         column("solid_mole_fraction", time=2., solid_species="NiO")]))
    np.testing.assert_equal(list(table_rows(dataset, profile)),
                            np.column_stack([dataset.x_cell, dataset.temperature.sel(time=1.),
                                             dataset.solid_mole_fraction.sel(time=2., solid_species="NiO")]))
    delta = plan_sheet(schema, sheet("delta", [column("distribution", time=1., x_cell=.45, gas_species="CO")]))
    assert delta["headings"][0] == "Oxygen deficiency (1)"
    np.testing.assert_equal(np.asarray(list(table_rows(dataset, delta)))[:, 1],
                            dataset.distribution.sel(time=1., x_cell=.45, gas_species="CO"))
    categorical = plan_sheet(schema, sheet("gas_species", [column("gas_mole_fraction", time=1., x_cell=.15)]))
    assert [r[0] for r in table_rows(dataset, categorical)] == ["CH4", "CO", "H2"]


def test_selection_failures_and_roundoff(dataset, monkeypatch):
    schema = describe_dataset(dataset)
    for bad, message in [
        (sheet("x_cell", [column("velocity", time=0.)]), "does not have"),
        (sheet(columns=[column("temperature", x_cell=.16)]), "not recorded"),
        (sheet(columns=[column("gas_mole_fraction", x_cell=.15)]), "every other axis"),
        (sheet(columns=[column("temperature", x_cell=.15)] * 2), "Duplicate heading"),
        (sheet(columns=[column("pressure", x_cell=.15)]), "General"),
    ]:
        with pytest.raises(ValueError, match=message):
            plan_sheet(schema, bad)
    assert coordinate_index(np.array([.3]), .1 + .2, "time") == 0
    with pytest.raises(ValueError, match="ambiguous"):
        coordinate_index(np.array([1., 1.]), 1., "time")
    unknown = dataset.drop_vars("delta")
    with pytest.raises(ValueError, match="one-dimensional coordinate"):
        plan_sheet(describe_dataset(unknown), sheet("delta", [column("distribution", time=0., x_cell=.15, gas_species="CH4")]))
    for name in ("Case information", "History", "a/b", "'name", "x" * 32):
        with pytest.raises(ValueError, match="sheet name"):
            plan_workbook(schema, report(sheet(columns=[column("outlet_temperature")], name=name)))
    monkeypatch.setattr(workbook, "MAX_ROWS", 3)
    with pytest.raises(ValueError, match="limits"):
        plan_sheet(schema, sheet(columns=[column("outlet_temperature")]))


def test_row_order_endpoints_and_missing_samples(dataset):
    schema = describe_dataset(dataset)
    col = {"quantity": "temperature", "fixed": {"x_cell": {"mode": "last"}}}
    for mode, expected in (("all", [0., 1., 2.]), ("first", [0.]), ("last", [2.])):
        table = plan_sheet(schema, sheet(columns=[col], rows={"mode": mode}))
        assert table["values"].tolist() == expected
        assert table["columns"][0]["fixed"] == {"x_cell": .75}
    table = plan_sheet(schema, sheet(columns=[col], rows={"mode": "values", "values": [2., 0.]}))
    assert [r[0] for r in table_rows(dataset, table)] == [2., 0.]
    partial = describe_dataset(dataset.isel(time=slice(0, 2)))
    assert len(plan_sheet(partial, sheet(columns=[col]))["rows"]) == 2
    with pytest.raises(ValueError, match="not recorded"):
        plan_sheet(partial, sheet(columns=[col], rows={"mode": "values", "values": [2.]}))


def test_expected_schema_and_domain_discovery(recorded, monkeypatch):
    case = recorded
    schema = describe_inputs(case.documents)
    np.testing.assert_equal(schema["axes"]["time"]["values"], [.0, .01])
    np.testing.assert_allclose(schema["axes"]["x_cell"]["values"], [1/6, .5, 5/6])
    assert schema["quantities"]["outlet_temperature"]["unit"] == "K"
    np.testing.assert_allclose(reporting_times(1.1, .5), [0., .5, 1., 1.1])
    bad = deepcopy(case.documents)
    bad["run"]["model"]["bed_length_m"] = ""
    bad["run"]["simulation"]["reporting_interval_s"] = ""
    schema = describe_inputs(bad)
    assert schema["axes"]["x_cell"]["values"] is None
    assert schema["axes"]["time"]["values"] is None
    assert schema["axes"]["gas_species"]["values"].tolist() == ["N2"]

    from packed_bed import report_schema
    monkeypatch.setattr(report_schema, "REPORT_REGISTRY", {
        "future": ReportSpec("Future quantity", (ModelField("source", "distribution", ("delta",)),), unit="mol",
                             axis_resolvers={"delta": lambda _: {"values": [.3], "units": "1"}}),
    })
    documents = {"run": {"outputs": {"requested_reports": ["future"]}}}
    variable = SimpleNamespace(Domains=[SimpleNamespace(Points=[.1, .2], Units="1")])
    schema = describe_inputs(documents, variables={"source": variable})
    assert schema["axes"]["delta"]["values"].tolist() == [.1, .2]
    fallback = describe_inputs(documents)
    assert fallback["axes"]["delta"]["values"].tolist() == [.3]


def test_excel_tables_dictionary_and_retained_provenance(recorded, dataset, tmp_path):
    case = recorded
    original = case.documents["run"]["model"]["bed_length_m"]
    case.documents["run"]["model"]["bed_length_m"] = 42
    cols = [column("gas_mole_fraction", gas_species="CH4", x_cell=.15),
            {**column("outlet_temperature"), "label": "=literal"}]
    definition = report(sheet(columns=cols, name="History data"),
                        sheet("x_cell", [column("temperature", time=1.)], name="Profile"))
    destination = tmp_path / "report.xlsx"
    write_workbook(case.run_folder, definition, destination, fingerprint=case.fingerprint())
    book = load_workbook(destination)
    assert book.sheetnames == ["Case information", "History data", "Profile"]
    data = book["History data"]
    assert data.freeze_panes == "B2"
    assert data.tables["Report1"].ref == "A1:C4"
    assert data.tables["Report1"].autoFilter.ref == "A1:C4"
    assert [c.name for c in data.tables["Report1"].tableColumns] == [c.value for c in data[1]]
    assert data["C1"].value == "=literal (K)" and data["C1"].data_type == "s"
    np.testing.assert_equal(list(data.values)[1:], list(table_rows(dataset, plan_sheet(describe_dataset(dataset), definition["sheets"][0]))))
    rows = list(book["Case information"].values)
    assert any(r[:3] == ("Inputs", "run.model.bed_length_m", original) for r in rows)
    assert any(r[:3] == ("Export", "inputs_changed", True) for r in rows)
    assert any("retained-hash" in r for r in rows)
    assert any(r[0] == "History data" and .15 in r for r in rows)
    book.close()


def test_atomic_failure_cancellation_and_nonfinite(recorded, tmp_path, monkeypatch):
    definition = report(sheet(columns=[column("outlet_temperature")]))
    destination = tmp_path / "existing.xlsx"
    destination.write_bytes(b"previous workbook")
    with pytest.raises(ExportCancelled):
        write_workbook(recorded.run_folder, definition, destination, cancelled=lambda: True)
    assert destination.read_bytes() == b"previous workbook"
    def interrupted(*args, **kwargs):
        yield [0., 300.]
        raise ExportCancelled()
    with monkeypatch.context() as patch:
        patch.setattr(workbook, "table_rows", interrupted)
        with pytest.raises(ExportCancelled):
            write_workbook(recorded.run_folder, definition, destination)
    assert destination.read_bytes() == b"previous workbook"
    replace = Path.replace
    def locked(path, target):
        if Path(target) == destination:
            raise PermissionError("File locked by Excel")
        return replace(path, target)
    with monkeypatch.context() as patch:
        patch.setattr(Path, "replace", locked)
        with pytest.raises(PermissionError, match="locked"):
            write_workbook(recorded.run_folder, definition, destination)
    assert destination.read_bytes() == b"previous workbook"
    assert not list(tmp_path.glob(".report-*.xlsx"))
    assert (recorded.run_folder / "output/results.nc").exists()
    assert workbook.excel_value(float("nan")) == "nan"
    assert workbook.excel_value(float("inf")) == "inf"
    with pytest.raises(ValueError, match="overlong"):
        workbook.excel_value("x" * 32768)


def test_editor_report_persistence_readonly_and_duplicate(qt_app, recorded, monkeypatch):
    from packed_bed_ui.editor import CaseEditor
    from packed_bed_ui.report import ColumnsDialog
    editor = CaseEditor()
    assert editor.set_case(recorded)
    page = editor.report
    fingerprint = recorded.fingerprint()
    inputs = deepcopy(recorded.documents)
    assert page.sheets.currentRow() == 1
    assert page.sheet["columns"] == []
    def accept(dialog):
        dialog.quantity.setCurrentIndex(dialog.quantity.findData("gas_mole_fraction"))
        assert dialog.count == 3  # Three cells and one species in the current inputs.
        return dialog.DialogCode.Accepted
    monkeypatch.setattr(ColumnsDialog, "exec", accept)
    page.add_columns()
    assert len(page.sheet["columns"]) == 3
    def add_duplicate(dialog):
        dialog.quantity.setCurrentIndex(dialog.quantity.findData("gas_mole_fraction"))
        for picker in dialog.pickers.values():
            picker.set_selection({"mode": "values", "values": [workbook.scalar(picker.axis["values"][0])]})
        return dialog.DialogCode.Accepted
    monkeypatch.setattr(ColumnsDialog, "exec", add_duplicate)
    page.add_columns()
    assert len(page.sheet["columns"]) == 3
    page.columns.setCurrentRow(0)
    page.remove_columns()
    page.columns.setCurrentRow(0)
    page.move_column(1)
    assert len(page.sheet["columns"]) == 2
    assert editor.save()
    assert recorded.fingerprint() == fingerprint
    assert recorded.documents == inputs
    reopened = Project.open(recorded.project.root).cases[0]
    assert reopened.metadata["report"] == page.definition
    duplicate = recorded.project.duplicate_case(recorded, "Copy")
    duplicate.metadata["report"]["sheets"][0]["name"] = "Independent"
    assert recorded.metadata["report"]["sheets"][0]["name"] != "Independent"
    assert not duplicate.run_folder.exists()
    recorded.metadata["study_id"] = "generated"
    assert editor.set_case(recorded)
    assert editor.read_only
    page.add_sheet()
    assert editor.save()
    assert len(recorded.metadata["report"]["sheets"]) == 2
    editor.close()


def test_report_save_failure_blocks_switch_and_preserves_draft(qt_app, recorded, monkeypatch):
    from packed_bed_ui.editor import CaseEditor
    editor = CaseEditor()
    editor.set_case(recorded)
    editor.report.add_sheet()
    other = recorded.project.add_case("Other")
    def failure():
        raise OSError("read only")
    with monkeypatch.context() as patch:
        patch.setattr(recorded.project, "save", failure)
        assert not editor.set_case(other)
        assert editor.report.case is recorded and editor.report.dirty
        assert "report" not in recorded.metadata
    assert editor.save()
    assert editor.set_case(other)
    assert editor.report.sheets.currentRow() == 1
    assert editor.report.sheet["columns"] == []
    assert "report" not in other.metadata
    editor.close()


def test_coordinate_labels_keep_distinct_values(dataset):
    assert workbook.text_value(.5249999999999999) == "0.525"
    values = np.array([.5, np.nextafter(.5, 1.)])
    close = dataset.isel(x_cell=slice(0, 2)).assign_coords(x_cell=values)
    table = plan_sheet(describe_dataset(close), sheet(columns=[column("temperature", x_cell=v) for v in values]))
    assert table["headings"][1] != table["headings"][2]


def test_coordinate_picker_preserves_selections_when_filtering(qt_app, monkeypatch):
    from PyQt6.QtCore import QItemSelectionModel
    from PyQt6.QtWidgets import QDialog, QLineEdit, QListView
    from packed_bed_ui.report import CoordinatePicker
    picker = CoordinatePicker(axis_description("x_cell", [.15, .45, .75]))
    def choose(dialog):
        view = dialog.findChild(QListView)
        search = dialog.findChild(QLineEdit)
        view.selectionModel().select(view.model().index(0, 0), QItemSelectionModel.SelectionFlag.Select)
        search.setText("0.75")
        view.selectionModel().select(view.model().index(0, 0), QItemSelectionModel.SelectionFlag.Select)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(QDialog, "exec", choose)
    picker.choose_values()
    assert picker.selection() == {"mode": "values", "values": [.15, .75]}
    assert picker.options() == [{"value": .15}, {"value": .75}]
    picker.close()


def test_recorded_preview_and_never_run_configuration(qt_app, recorded, source_case):
    from packed_bed_ui.editor import CaseEditor
    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    page.sheet["columns"] = [column("outlet_temperature")]
    page.select_sheet()
    assert not hasattr(page, "source")
    assert [[page.table.item(i, j).text() for j in range(2)] for i in range(3)] == [
        ["0", "301"], ["1", "304"], ["2", "307"]]
    assert "Recorded results · 3 rows" in page.counts.text()
    recorded.documents["run"]["model"]["bed_length_m"] = ""
    page.refresh_source()
    assert page.table.item(0, 1).text() == "301"
    assert "Recorded results · 3 rows" in page.counts.text()
    assert page.schema["axes"]["x_cell"]["values"] is None
    assert page.export_button.isEnabled()
    never_run = recorded.project.add_case_from_files(source_case, "Future")
    # Flush the existing case before checking that opening the next one is read-only.
    assert editor.save()
    project_bytes = (recorded.project.root / "project.json").read_bytes()
    editor.set_case(never_run)
    assert (recorded.project.root / "project.json").read_bytes() == project_bytes
    assert "report" not in never_run.metadata
    assert page.sheets.currentRow() == 1
    page.sheet["columns"] = [{"quantity": "temperature", "fixed": {"x_cell": {"mode": "last"}}}]
    page.select_sheet()
    assert "Expected" in page.counts.text() and not page.export_button.isEnabled()
    assert "No recorded results" in page.counts.text()
    assert page.table.item(0, 1).text() == ""
    assert editor.save()
    editor.close()


def test_preview_resolves_recorded_coordinates_and_missing_values(qt_app, recorded, dataset):
    from packed_bed_ui.editor import CaseEditor
    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    page.sheet.update(sheet("x_cell", [
        {"quantity": "temperature", "fixed": {"time": {"mode": "last"}}},
    ], rows={"mode": "values", "values": [.75, .15]}))
    definition = deepcopy(page.definition)
    page.select_sheet()
    assert [[page.table.item(i, j).text() for j in range(2)] for i in range(2)] == [
        ["0.75", "308"], ["0.15", "306"]]
    assert "Time=2 s" in page.table.horizontalHeaderItem(1).toolTip()
    assert page.definition == definition

    # A replacement/partial run must refresh the values and resolve Last again.
    dataset.isel(time=slice(0, 2)).to_netcdf(recorded.run_folder / "output/results.nc", engine="scipy")
    page.preview()
    assert page.table.item(0, 1).text() == "305"
    assert "Time=1 s" in page.table.horizontalHeaderItem(1).toolTip()
    assert page.definition == definition

    page.sheet["columns"][0]["fixed"]["time"] = {"value": 2.}
    page.preview()
    assert "not recorded" in page.error.text()
    assert page.table.rowCount() == 0
    assert not page.export_button.isEnabled()

    page.sheet.update(sheet(columns=[column("pressure", x_cell=.15)]))
    page.preview()
    assert "Quantity 'pressure' is unavailable" in page.error.text()
    assert not page.export_button.isEnabled()
    editor.close()


def test_preview_matches_workbook_with_bounded_values_and_run_information(qt_app, recorded, dataset, tmp_path):
    from packed_bed_ui.editor import CaseEditor
    extended = xr.concat([dataset] * 10, dim="time").assign_coords(time=np.arange(30.))
    extended.to_netcdf(recorded.run_folder / "output/results.nc", engine="scipy")
    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    page.sheet["columns"] = expand_columns("distribution", {
        "x_cell": [{"value": v} for v in [.15, .45, .75]],
        "gas_species": [{"value": v} for v in ["CH4", "CO", "H2"]],
        "delta": [{"value": v} for v in [0., .2]],
    })
    page.select_sheet()
    assert page.table.rowCount() == 20
    assert page.table.columnCount() == 12
    assert "30 rows × 19 columns" in page.counts.text()
    destination = tmp_path / "preview.xlsx"
    write_workbook(recorded.run_folder, page.definition, destination)
    book = load_workbook(destination)
    for i in range(20):
        for j in range(12):
            assert float(page.table.item(i, j).text()) == book["Sheet 1"].cell(i + 2, j + 1).value
    book.close()

    page.sheets.setCurrentRow(0)
    assert "Recorded run information" in page.counts.text()
    assert any(page.table.item(i, 2).text() == "attempt-one" for i in range(page.table.rowCount()))
    page.info_view.setCurrentIndex(1)
    assert page.table.item(0, 7).text() == "30"
    assert not page.error.text()
    editor.close()


def test_export_dialog_worker_and_cancellation(qt_app, recorded, tmp_path):
    from PyQt6.QtCore import QTimer
    from packed_bed_ui.report import ExportDialog
    destination = tmp_path / "worker.xlsx"
    arguments = {"run_folder": recorded.run_folder,
                 "definition": report(sheet(columns=[column("outlet_temperature")])),
                 "destination": destination}
    for cancel in (False, True):
        destination.write_bytes(b"existing")
        dialog = ExportDialog(arguments)
        if cancel:
            dialog.reject()  # Also works before QThread.start() is delivered.
        timeout = QTimer(dialog)
        timeout.setSingleShot(True)
        timeout.timeout.connect(dialog.reject)
        timeout.start(10000)
        dialog.exec()
        assert dialog.worker.wait(1000)
        assert not dialog.worker.error
        assert dialog.worker.cancelled == cancel
        if cancel:
            assert destination.read_bytes() == b"existing"
        else:
            book = load_workbook(destination)
            assert book["Data"]["B2"].value == 301
            book.close()
        dialog.deleteLater()


def test_all_columns_follow_inputs_preserving_edits_and_export_values(qt_app, recorded, monkeypatch, tmp_path):
    from packed_bed_ui.editor import CaseEditor
    from packed_bed_ui.report import ColumnsDialog

    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    def accept(dialog):
        dialog.quantity.setCurrentIndex(dialog.quantity.findData("temperature"))
        return dialog.DialogCode.Accepted
    monkeypatch.setattr(ColumnsDialog, "exec", accept)
    page.add_columns()
    page.columns.setCurrentRow(1)
    page.remove_columns()
    page.sheet["columns"][1]["label"] = "Final original cell"
    page.columns.setCurrentRow(1)
    page.move_column(-1)
    assert editor.save()
    # Geometry and grid changes are made on the Bed/General pages, with Report inactive.
    editor.put(("run", "model", "bed_length_m"), 2.)
    editor.put(("run", "model", "axial_cells"), 4)
    assert editor.save()
    table = plan_sheet(page.schema, page.sheet)
    assert [c["fixed"]["x_cell"] for c in table["columns"]] == [1.25, .25, 1.75]
    assert table["headings"][1] == "Final original cell (K)"
    assert len(page.sheet["columns"]) == 3  # Removed cell 2 stays removed; new cell 4 appears.
    assert page.columns.count() == 3
    # Export resolves the same All rule against the retained grid, using actual recorded values.
    original = deepcopy(page.definition)
    destination = tmp_path / "adapted.xlsx"
    write_workbook(recorded.run_folder, page.definition, destination)
    book = load_workbook(destination)
    assert book["Sheet 1"]["B2"].value == 302
    assert book["Sheet 1"]["C2"].value == 300
    assert book["Sheet 1"].max_column == 3
    page.preview()
    assert page.table.item(0, 1).text() == "302"
    assert page.table.item(0, 2).text() == "300"
    assert page.table.columnCount() == 3
    assert page.definition == original
    editor.close()


def test_exact_coordinates_are_removed_only_after_inputs_define_a_new_axis(qt_app, recorded):
    from packed_bed_ui.editor import CaseEditor

    recorded.metadata["report"] = report(
        sheet(columns=[column("temperature", x_cell=.5),
                       {"quantity": "temperature", "fixed": {"x_cell": {"mode": "last"}}}]),
        sheet("x_cell", [column("temperature", time=0.)], "Profile", {"mode": "values", "values": [.5]}))
    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    editor.put(("run", "model", "bed_length_m"), "")
    assert editor.save()
    assert len(page.sheet["columns"]) == 2  # A temporarily invalid draft doesn't discard selections.
    editor.put(("run", "model", "bed_length_m"), 2.)
    assert editor.save()
    assert len(page.sheet["columns"]) == 1
    assert page.sheet["columns"][0]["fixed"] == {"x_cell": {"mode": "last"}}
    assert page.definition["sheets"][1]["rows"] == {"mode": "values", "values": []}
    assert "Removed 2" in page.axis_notice.text()
    reopened = Project.open(recorded.project.root).cases[0]
    assert reopened.metadata["report"] == page.definition
    editor.close()


def test_templates_are_portable_independent_and_reviewed_before_apply(qt_app, recorded, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QFileDialog, QDialog, QTableWidget
    from packed_bed_ui.editor import CaseEditor
    from packed_bed_ui.report import ColumnsDialog, TemplateDialog
    from packed_bed_ui.project import read_json

    editor = CaseEditor()
    editor.set_case(recorded)
    page = editor.report
    def columns(dialog):
        dialog.quantity.setCurrentIndex(dialog.quantity.findData("temperature"))
        return dialog.DialogCode.Accepted
    monkeypatch.setattr(ColumnsDialog, "exec", columns)
    page.add_columns()
    page.rename_sheet("Temperatures")
    path = tmp_path / "template.json"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(path), ""))
    page.save_template()
    template_bytes = path.read_bytes()
    saved = read_json(path)
    assert set(saved) == {"version", "sheets"}

    other = recorded.project.duplicate_case(recorded, "Other")
    other.documents["run"]["model"]["axial_cells"] = 5
    other.documents["run"]["model"]["bed_length_m"] = 2.
    other.metadata.pop("report")
    other.save()
    editor.set_case(other)
    previous = deepcopy(page.definition)
    fingerprint = other.fingerprint()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(path), ""))
    def review(dialog):
        table = dialog.findChild(QTableWidget)
        assert table.item(0, 0).text() == "Temperatures"
        assert table.item(0, 2).text() == "2 rows × 6 columns"
        assert table.item(0, 3).text() == "Ready"
        return QDialog.DialogCode.Rejected
    monkeypatch.setattr(TemplateDialog, "exec", review)
    page.apply_template()
    assert page.definition == previous and "report" not in other.metadata
    monkeypatch.setattr(TemplateDialog, "exec", lambda dialog: QDialog.DialogCode.Accepted)
    page.apply_template()
    assert len(page.sheet["columns"]) == 5
    page.rename_sheet("Adapted")
    assert editor.save()
    assert other.fingerprint() == fingerprint
    assert path.read_bytes() == template_bytes
    assert recorded.metadata["report"]["sheets"][0]["name"] == "Temperatures"
    assert Project.open(recorded.project.root).cases[1].metadata["report"]["sheets"][0]["name"] == "Adapted"

    # Malformed/unsupported files leave the current report untouched.
    previous = deepcopy(page.definition)
    path.write_text('{"version": 99, "sheets": []}')
    page.apply_template()
    assert page.definition == previous
    assert "Unsupported" in page.error.text()
    editor.close()


def test_all_rules_respect_column_limits_before_expanding(recorded, monkeypatch):
    schema = describe_inputs(recorded.documents)
    rule = {"id": "temperature", "quantity": "temperature", "selections": {"x_cell": {"mode": "all"}}}
    monkeypatch.setattr(workbook, "MAX_COLUMNS", 3)
    with pytest.raises(ValueError, match="too many"):
        workbook.rule_columns(schema, rule)


def test_template_preserves_unavailable_quantities_as_repairable_errors(qt_app, recorded):
    from PyQt6.QtWidgets import QTableWidget
    from packed_bed_ui.report import TemplateDialog

    definition = report(sheet(columns=[column("not_recorded")]))
    dialog = TemplateDialog(describe_inputs(recorded.documents), definition)
    assert "unavailable" in dialog.findChild(QTableWidget).item(0, 3).text()
    assert definition["sheets"][0]["columns"] == [column("not_recorded")]
    dialog.close()


@pytest.mark.parametrize("bad", [
    report(sheet(rows={"mode": "values", "values": None})),
    report(sheet(rows={"mode": "values", "values": [[1]]})),
    report(sheet(columns=[column("temperature", x_cell=[1, 2])])),
])
def test_malformed_template_selectors_are_rejected(bad):
    with pytest.raises(ValueError, match="Malformed"):
        workbook.report_definition(bad)
