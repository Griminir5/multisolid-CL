"""Project exports preserve each run's recorded values, grids and provenance."""

from copy import deepcopy

import numpy as np
import pytest
import xarray as xr
from openpyxl import load_workbook

from packed_bed_ui.project import Project, write_documents, write_json
from packed_bed_ui.project_results import (
    open_results, plan_project_sheet, plan_project_workbook, project_columns,
    project_information_rows, project_table_rows, split_by_case, write_project_workbook,
)
from packed_bed_ui.workbook import ExportCancelled


@pytest.fixture
def runs(tmp_path, source_case):
    project = Project.create(tmp_path / "project")
    def add(name, times=(0., 1., 2.), positions=(.25, .75), offset=0., missing=False, state="completed"):
        case = project.add_case_from_files(source_case, name)
        write_documents(case.run_folder / "inputs", case.documents)
        write_json(case.run_folder / "snapshot.json", {
            "case_id": case.id, "case_name": name, "attempt_id": f"attempt-{case.id}",
            "fingerprint": case.fingerprint(),
        })
        write_json(case.run_folder / "status.json", {"state": state})
        (case.run_folder / "output").mkdir()
        ds = xr.Dataset({
            "outlet_composition": (("time", "gas_species"),
                                   np.arange(len(times) * 3).reshape(-1, 3) + offset, {"units": "1"}),
            "temperature": (("time", "x_cell"),
                            np.arange(len(times) * len(positions)).reshape(len(times), -1) + 300. + offset,
                            {"units": "K"}),
        }, coords={"time": ("time", list(times), {"units": "s"}), "gas_species": ["H2", "CO", "CO2"],
                   "x_cell": ("x_cell", list(positions), {"units": "m"})})
        if missing:
            ds = ds.drop_vars("outlet_composition")
        ds.to_netcdf(case.run_folder / "output/results.nc", engine="scipy")
        return case, ds
    return project, add


def sheet(columns, axis="time", name="Data", rows=None):
    return {"name": name, "axis": axis, "columns": columns, "rows": rows or {"mode": "all"}}


def report(cases, *sheets):
    return {"version": 1, "case_ids": [c.id for c in cases], "sheets": list(sheets)}


def outlet(case, species="H2"):
    return {"quantity": "outlet_composition", "fixed": {"case": {"value": case.id}, "gas_species": {"value": species}}}


def test_five_cases_three_components_match_excel(runs, tmp_path):
    project, add = runs
    cases, datasets = zip(*(add(f"Case {i}", offset=100 * i) for i in range(5)))
    with open_results(project, [c.id for c in cases]) as schema:
        columns = project_columns(schema, "outlet_composition", {
            "case": {"mode": "all"}, "gas_species": {"mode": "all"}})
        assert len(columns) == 15
        definition = report(cases, sheet(columns))
        table, = plan_project_workbook(schema, definition)
        expected = np.column_stack([datasets[0].time, *[ds.outlet_composition for ds in datasets]])
        np.testing.assert_equal(list(project_table_rows(schema, table)), expected)
    path = tmp_path / "results.xlsx"
    write_project_workbook(project, definition, path)
    book = load_workbook(path)
    assert book.sheetnames == ["Results information", "Data"]
    assert book["Data"].freeze_panes == "B2"
    assert book["Data"].tables["Report1"].ref == "A1:P4"
    np.testing.assert_equal(list(book["Data"].values)[1:], expected)
    headers = list(book["Data"].values)[0]
    assert "Case=Case 4" in headers[-1] and "CO2" in headers[-1]
    book.close()


def test_case_rows_resolve_last_per_run_and_document_actual_coordinates(runs):
    project, add = runs
    first, a = add("First")
    second, b = add("Second", times=(0., .5), positions=(.1, .4, .9), offset=10.)
    column = {"quantity": "temperature", "fixed": {"time": {"mode": "last"}, "x_cell": {"mode": "last"}}}
    with open_results(project, [first.id, second.id]) as schema:
        table = plan_project_sheet(schema, sheet([column], axis="case"))
        assert list(project_table_rows(schema, table)) == [["First", 305.], ["Second", 315.]]
        dictionary = list(project_information_rows(schema, [table]))
        assert any('"Time (s)": 2.0' in str(row[-1]) and '"Cell position (m)": 0.75' in str(row[-1]) for row in dictionary if row)
        assert any('"Time (s)": 0.5' in str(row[-1]) and '"Cell position (m)": 0.9' in str(row[-1]) for row in dictionary if row)


def test_different_time_grids_split_without_interpolation(runs, tmp_path):
    project, add = runs
    a, da = add("A")
    b, db = add("B", times=(0., .5, 1.), offset=10.)
    combined = sheet([outlet(a), outlet(b)])
    with open_results(project, [a.id, b.id]) as schema:
        with pytest.raises(ValueError, match="Split by case"):
            plan_project_sheet(schema, combined)
        sheets = split_by_case(schema, combined, ["Data"])
        tables = plan_project_workbook(schema, report([a, b], *sheets))
        for table, ds in zip(tables, [da, db]):
            np.testing.assert_equal(list(project_table_rows(schema, table)),
                                    np.column_stack([ds.time, ds.outlet_composition.sel(gas_species="H2")]))
        common = {**combined, "rows": {"mode": "values", "values": [0., 1.]}}
        table = plan_project_sheet(schema, common)
        assert list(project_table_rows(schema, table)) == [[0., 0., 10.], [1., 3., 16.]]
    destination = tmp_path / "different.xlsx"
    write_project_workbook(project, report([a, b], *sheets), destination)
    book = load_workbook(destination)
    assert list(book[sheets[1]["name"]].values)[2][0] == .5
    book.close()


def test_all_positions_expand_native_grids_and_exact_selection_never_snaps(runs):
    project, add = runs
    a, da = add("A")
    b, db = add("B", positions=(.1, .4, .9))
    with open_results(project, [a.id, b.id]) as schema:
        columns = project_columns(schema, "temperature", {"case": {"mode": "all"}, "x_cell": {"mode": "all"}})
        assert len(columns) == 5  # 2 + 3, never 2 cases × a union of 5 positions.
        table = plan_project_sheet(schema, sheet(columns))
        np.testing.assert_equal(list(project_table_rows(schema, table)), np.column_stack([da.time, da.temperature, db.temperature]))
        columns[0]["fixed"]["x_cell"] = {"value": np.nextafter(.25, 1.)}
        with pytest.raises(ValueError, match="not uniquely recorded"):
            plan_project_sheet(schema, sheet(columns))


def test_missing_results_quantities_units_and_deleted_cases_are_explicit(runs):
    project, add = runs
    a, _ = add("Available")
    b, _ = add("Not recorded", missing=True)
    c = project.add_case("Not run")
    with open_results(project, [a.id, b.id, c.id, "deleted"]) as schema:
        for case, message in [(b, "was not recorded"), (c, "No readable" )]:
            with pytest.raises(ValueError, match=message):
                plan_project_sheet(schema, sheet([outlet(case)]))
        with pytest.raises(ValueError, match="deleted or replaced"):
            plan_project_sheet(schema, sheet([{**outlet(a), "fixed": {"case": {"value": "deleted"}, "gas_species": {"value": "H2"}}}]))
        assert schema["sources"][c.id]["error"]
        # Unselected/unreferenced failures do not prevent exporting a valid sheet.
        assert len(plan_project_sheet(schema, sheet([outlet(a)]))["rows"]) == 3
        schema["sources"][a.id]["schema"]["quantities"]["outlet_composition"]["unit"] = "mol/s"
        with pytest.raises(ValueError, match="incompatible units"):
            plan_project_sheet(schema, sheet([outlet(a)]))


def test_stale_export_retains_snapshot_and_atomic_failure(runs, tmp_path):
    project, add = runs
    a, ds = add("=Literal case", times=(0., .5))
    a.documents["run"]["model"]["bed_length_m"] = 99
    definition = report([a], sheet([outlet(a)]))
    path = tmp_path / "data.xlsx"
    write_project_workbook(project, definition, path)
    book = load_workbook(path)
    rows = list(book["Results information"].values)
    assert any(row[:3] == (a.name, "Export.inputs_changed", True) for row in rows)
    assert any(row[:3] == (a.name, "Inputs.run.model.bed_length_m", 1) for row in rows)
    assert any(row[:3] == (a.name, "Status.state", "completed") for row in rows)
    assert book["Results information"]["A4"].data_type == "s"
    assert book["Data"].max_row == 3
    book.close()
    previous = path.read_bytes()
    with pytest.raises(ExportCancelled):
        write_project_workbook(project, definition, path, cancelled=lambda: True)
    assert path.read_bytes() == previous
    definition["sheets"][0]["columns"][0]["fixed"]["gas_species"] = {"value": "Missing"}
    with pytest.raises(ValueError, match="not uniquely recorded"):
        write_project_workbook(project, definition, path)
    assert path.read_bytes() == previous
    assert not list(tmp_path.glob(".report-*.xlsx"))


def test_results_ui_reuses_report_and_saves_without_changing_inputs(qt_app, runs, monkeypatch, tmp_path):
    from PyQt6.QtWidgets import QDialog, QTreeWidget
    from PyQt6.QtCore import Qt
    from packed_bed_ui.results import ProjectColumnsDialog
    from packed_bed_ui.window import MainWindow

    project, add = runs
    a, _ = add("A")
    b, _ = add("B", times=(0., .5, 1.))
    fingerprints = [c.fingerprint() for c in project.cases]
    window = MainWindow()
    window._set_project(project)
    page = window.results
    def choose(dialog):
        assert window.pages.currentWidget() is window.home
        for item in dialog.findChild(QTreeWidget).findItems("", Qt.MatchFlag.MatchContains):
            item.setCheckState(0, Qt.CheckState.Checked)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(QDialog, "exec", choose)
    window.results_button.click()
    assert window.pages.currentWidget() is page
    assert page.definition["case_ids"] == [a.id, b.id]
    assert page.sheet["axis"] == "time"
    def columns(dialog):
        dialog.quantity.setCurrentIndex(dialog.quantity.findData("outlet_composition"))
        assert dialog.count == 6
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ProjectColumnsDialog, "exec", columns)
    page.add_columns()
    assert "Split by case" in page.error.text()
    assert len(page.sheet["columns"]) == 6
    assert page.rows.note.text() == "Per case"
    np.testing.assert_equal(page.rows.axis["values"], [0., 1.])
    page.split_sheet()
    assert len(page.definition["sheets"]) == 2
    assert page.table.rowCount() == 3 and page.table.columnCount() == 4
    assert page.rows.note.text() == "3 values"
    np.testing.assert_equal(page.rows.axis["values"], [0., 1., 2.])
    assert page.export_button.isEnabled()
    page.sheets.setCurrentRow(2)
    assert page.table.item(1, 0).text() == "0.5"
    np.testing.assert_equal(page.rows.axis["values"], [0., .5, 1.])
    assert window._save_editors()
    assert [c.fingerprint() for c in project.cases] == fingerprints
    assert not any("report" in c.metadata for c in project.cases)
    assert Project.open(project.root).metadata["results_report"] == page.definition
    assert not project.drafts.pending()
    window.close()


def test_results_save_failure_blocks_navigation_and_recovers(qt_app, runs, monkeypatch):
    from packed_bed_ui.window import MainWindow
    from packed_bed_ui.results import CaseSelectionDialog
    from PyQt6.QtCore import Qt

    project, add = runs
    case, _ = add("A")
    window = MainWindow()
    window._set_project(project)
    def choose(dialog):
        dialog.items[case.id].setCheckState(0, Qt.CheckState.Checked)
        return dialog.DialogCode.Accepted
    monkeypatch.setattr(CaseSelectionDialog, "exec", choose)
    window._show_results()
    page = window.results
    page.definition = report([case], sheet([outlet(case)]))
    page.changed()
    expected = deepcopy(page.definition)
    def fail():
        raise OSError("read only")
    with monkeypatch.context() as patch:
        patch.setattr(project, "save", fail)
        window._show_cases()
        assert window.pages.currentWidget() is page
        assert page.dirty and "read only" in window.statusBar().currentMessage()
        assert "results_report" not in project.metadata
        recovered = Project.open(project.root)
        draft, = recovered.drafts.pending()
        recovered.drafts.apply(draft)
        assert recovered.metadata["results_report"] == expected
        assert recovered.cases[0].fingerprint() == case.fingerprint()
    assert window._save_editors()
    window._show_cases()
    assert window.pages.currentWidget() is window.home
    window.close()


def test_results_case_rows_preview_and_rerun_preserve_missing_selections(qt_app, runs):
    from packed_bed_ui.results import ResultsPage

    project, add = runs
    case, ds = add("A")
    definition = report([case], sheet([{"quantity": "temperature", "fixed": {
        "time": {"mode": "last"}, "x_cell": {"mode": "last"}}}], axis="case"))
    project.metadata["results_report"] = definition
    page = ResultsPage()
    page.set_project(project)
    assert page.table.item(0, 1).text() == "305"
    ds.isel(time=slice(0, 2)).to_netcdf(case.run_folder / "output/results.nc", engine="scipy")
    page.refresh_source()
    assert page.table.item(0, 1).text() == "303"
    page.sheet["columns"][0]["fixed"]["time"] = {"value": 2.}
    saved = deepcopy(page.definition)
    page.refresh_source()
    assert "not uniquely recorded" in page.error.text()
    assert not page.export_button.isEnabled()
    assert page.definition == saved
    project.delete_case(case)
    page.refresh_source()
    assert page.definition == saved
    assert "deleted or replaced" in page.axis_notice.text()
    page.close()


def test_results_entry_cancel_and_success_only_selection(qt_app, runs, monkeypatch):
    from PyQt6.QtCore import Qt
    from PyQt6.QtWidgets import QDialogButtonBox, QHBoxLayout, QLineEdit, QPushButton
    from packed_bed_ui.results import CaseSelectionDialog
    from packed_bed_ui.window import MainWindow

    project, add = runs
    a, _ = add("First success")
    b, _ = add("Hidden success")
    failed, _ = add("Failed", state="failed")
    unrun = project.add_case("Not run")
    missing, _ = add("Missing data")
    (missing.run_folder / "output/results.nc").unlink()
    for case in (a, failed):
        case.metadata.update(study_id="study", origin="Study")
    project.metadata["results_report"] = report([failed], sheet([outlet(failed)]))
    dialog = CaseSelectionDialog(project, [failed.id, unrun.id, missing.id])
    assert not dialog.selected_cases()
    ok = dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Ok)
    assert not ok.isEnabled()
    for case in (failed, unrun, missing):
        assert dialog.items[case.id].isDisabled()
    dialog.findChild(QLineEdit).setText("First success")
    buttons = {button.text(): button for button in dialog.findChildren(QPushButton)}
    buttons["Select all"].click()
    assert dialog.selected_cases() == [a.id, b.id]
    assert ok.isEnabled()
    buttons["Clear all"].click()
    assert not dialog.selected_cases()
    dialog.items[a.id].parent().setCheckState(0, Qt.CheckState.Checked)
    assert dialog.selected_cases() == [a.id]
    assert dialog.items[failed.id].checkState(0) != Qt.CheckState.Checked
    dialog.close()

    # Restore ordinary case metadata before opening the full project.
    for case in (a, failed):
        case.metadata.pop("study_id")
    window = MainWindow()
    window._set_project(project)
    execution = next(layout for layout in window.home.findChildren(QHBoxLayout)
                     if layout.indexOf(window.run_all_button) >= 0)
    assert execution.indexOf(window.results_button) == execution.indexOf(window.run_all_button) + 1
    previous = deepcopy(project.metadata)
    monkeypatch.setattr(CaseSelectionDialog, "exec", lambda self: self.DialogCode.Rejected)
    window.results_button.click()
    assert window.pages.currentWidget() is window.home
    assert project.metadata == previous
    window.close()


def test_failed_run_excluded_after_saved_selection_but_case_report_exports_partial(qt_app, runs, tmp_path):
    from packed_bed_ui.editor import CaseEditor
    from packed_bed_ui.results import ResultsPage
    from packed_bed_ui.workbook import write_workbook

    project, add = runs
    case, ds = add("Failed", times=(0., .5), state="failed")
    project.metadata["results_report"] = report([case], sheet([outlet(case)]))
    page = ResultsPage()
    page.set_project(project)
    assert not page.export_button.isEnabled()
    assert "successful run" in page.axis_notice.text()
    with pytest.raises(ValueError, match="successful run"):
        write_project_workbook(project, project.metadata["results_report"], tmp_path / "blocked.xlsx")
    page.close()

    editor = CaseEditor()
    editor.set_case(case)
    editor.report.sheet["columns"] = [{"quantity": "temperature", "fixed": {"x_cell": {"mode": "last"}}}]
    editor.report.select_sheet()
    assert editor.report.export_button.isEnabled()
    assert editor.report.table.rowCount() == 2
    path = tmp_path / "partial.xlsx"
    write_workbook(case.run_folder, editor.report.definition, path)
    book = load_workbook(path)
    np.testing.assert_equal(list(book["Sheet 1"].values)[1:], np.column_stack([ds.time, ds.temperature.isel(x_cell=-1)]))
    assert any(row[:3] == ("Status", "state", "failed") for row in book["Case information"].values)
    book.close()
    editor.close()


def test_results_templates_are_portable_reviewed_and_independent(qt_app, runs, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QFileDialog, QDialog, QPushButton, QTableWidget
    from packed_bed_ui.results import ResultsPage
    from packed_bed_ui.report import TemplateDialog
    from packed_bed_ui.project import read_json
    from packed_bed_ui.project_results import results_template, apply_results_template

    project, add = runs
    a, _ = add("Source A")
    b, _ = add("Source B")
    project.metadata["results_report"] = report([a, b], sheet([outlet(a), outlet(b)]),
        sheet([{"quantity": "temperature", "fixed": {"time": {"mode": "last"}, "x_cell": {"mode": "last"}}}],
              axis="case", name="Endpoints", rows={"mode": "values", "values": [b.id, a.id]}))
    page = ResultsPage()
    page.set_project(project)
    buttons = {b.text() for b in page.findChildren(QPushButton)}
    assert {"Save as template…", "Apply template…"} <= buttons
    path = tmp_path / "template.json"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(path), ""))
    page.save_template()
    original = path.read_bytes()
    assert a.id.encode() not in original and b.id.encode() not in original
    assert b"attempt-" not in original
    template = read_json(path)
    assert template["case_count"] == 2
    c, _ = add("Target C", offset=10)
    d, _ = add("Target D", offset=20)
    project.metadata["results_report"] = report([c, d], sheet([outlet(c)]))
    page.set_project(project)
    before = deepcopy(page.definition)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(path), ""))
    def review(dialog):
        assert dialog.findChild(QTableWidget).item(0, 2).text() == "3 rows × 3 columns"
        assert dialog.findChild(QTableWidget).item(0, 3).text() == "Ready"
        return QDialog.DialogCode.Rejected
    monkeypatch.setattr(TemplateDialog, "exec", review)
    page.apply_template()
    assert page.definition == before
    monkeypatch.setattr(TemplateDialog, "exec", lambda self: QDialog.DialogCode.Accepted)
    page.apply_template()
    assert page.definition["case_ids"] == [c.id, d.id]
    assert page.sheet["columns"][0]["fixed"]["case"]["value"] == c.id
    assert page.definition["sheets"][1]["rows"]["values"] == [d.id, c.id]
    assert page.table.item(0, 1).text() == "10"
    assert page.table.item(0, 2).text() == "20"
    page.rename_sheet("Changed")
    assert page.save()
    assert path.read_bytes() == original
    assert Project.open(project.root).metadata["results_report"] == page.definition
    with pytest.raises(ValueError, match="Select 2"):
        apply_results_template(template, [c.id])
    bad = deepcopy(template)
    bad["sheets"][0]["columns"][0]["fixed"]["case"] = {"index": 99}
    with pytest.raises(ValueError, match="invalid case position"):
        apply_results_template(bad, [c.id, d.id])
    assert results_template(apply_results_template(template, [c.id, d.id])) == template
    page.close()


def test_real_midrun_failure_keeps_samples_exportable_from_case_report(qt_app, source_case, tmp_path, monkeypatch):
    dae = pytest.importorskip("daetools.pyDAE")
    from packed_bed.simulation import PackedBedSimulation
    from packed_bed_ui.worker import activate_snapshot, run_snapshot
    from packed_bed_ui.editor import CaseEditor
    from packed_bed_ui.workbook import write_workbook
    from packed_bed_ui.project import read_json

    project = Project.create(tmp_path / "project")
    case = project.add_case_from_files(source_case)
    case.documents["run"]["simulation"].update(time_horizon_s=.02, reporting_interval_s=.005)
    case.documents["run"]["outputs"]["requested_reports"] = ["temperature", "gas_concentration"]
    def fail_during_integration(simulation):
        simulation.ReportData(simulation.CurrentTime)
        for time in (.005, .01):
            simulation.IntegrateUntilTime(time, dae.eDoNotStopAtDiscontinuity)
            simulation.ReportData(simulation.CurrentTime)
        raise RuntimeError("Controlled failure during integration")
    monkeypatch.setattr(PackedBedSimulation, "Run", fail_during_integration)
    activate_snapshot(case.prepare("partial-test"))
    assert run_snapshot(case.run_folder) == 1
    assert read_json(case.run_folder / "status.json")["state"] == "failed"
    assert read_json(case.run_folder / "output/manifest.json")["status"] == "failed"

    editor = CaseEditor()
    editor.set_case(case)
    editor.report.sheet["columns"] = [
        {"quantity": "temperature", "fixed": {"x_cell": {"mode": "last"}}},
        {"quantity": "gas_concentration", "fixed": {"gas_species": {"value": "N2"}, "x_cell": {"mode": "last"}}},
    ]
    editor.report.select_sheet()
    assert editor.report.export_button.isEnabled()
    assert editor.report.table.rowCount() == 3
    destination = tmp_path / "partial.xlsx"
    write_workbook(case.run_folder, editor.report.definition, destination)
    book = load_workbook(destination)
    assert [row[0] for row in list(book["Sheet 1"].values)[1:]] == [0., .005, .01]
    assert book["Sheet 1"].max_column == 3
    book.close()
    editor.close()
