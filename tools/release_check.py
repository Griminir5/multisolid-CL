"""Installed-artifact acceptance check, invoked explicitly by release builders."""
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import traceback


def _command(*args):
    return [sys.executable, *([] if getattr(sys, "frozen", False) else ["-m", "packed_bed_ui"]), *map(str, args)]


def _documents():
    return {
        "run": {
            "references": {f"{key}_file": f"{key}.yaml" for key in ("chemistry", "program", "solids")},
            "simulation": {"system_name": "Release_check", "time_horizon_s": .1, "reporting_interval_s": .05,
                           "mass_scheme": "weno3", "heat_scheme": "weno3", "report_time_derivatives": False},
            "model": {"bed_length_m": 1., "bed_radius_m": .01, "axial_cells": 3,
                      "ambient_temperature_k": 300., "heat_transfer_coefficient_w_per_m2_k": 0.},
            "solver": {"backend": "daetools", "name": "superlu", "threads": 1, "relative_tolerance": 1e-5},
            "outputs": {"directory": "output", "artifacts_directory": "output/artifacts",
                        "requested_reports": ["temperature", "pressure", "gas_mole_fraction", "gas_flux"], "requested_plots": []}},
        "chemistry": {"gas_species": ["N2"], "reaction_families": [], "reaction_ids": []},
        "program": {"inlet_flow": {"initial": 1e-8, "steps": []},
                    "inlet_temperature": {"initial": 300., "steps": []},
                    "outlet_pressure": {"initial": 100000., "steps": []},
                    "inlet_composition": {"initial": {"N2": 1.}, "steps": []}},
        "solids": {"solid_species": ["Ni"], "initial_profile": {"basis": "solid", "zones": [
            {"x_start_m": 0., "x_end_m": 1., "e_b": .4, "e_p": .5, "d_p": .001, "values": {"Ni": 1.}}]}},
    }


def check(destination):
    destination = Path(destination).resolve()
    destination.mkdir(parents=True, exist_ok=False)
    record = {"executable": sys.executable, "frozen": bool(getattr(sys, "frozen", False)), "checks": []}

    def passed(name):
        record["checks"].append(name)
        (destination / "result.json").write_text(json.dumps(record, indent=2) + "\n")

    try:
        from packed_bed_ui.worker import diagnostic_log
        with diagnostic_log(destination / "self-test.log"):
            import numpy as np
            from packed_bed.simulation import create_linear_solver
            from packed_bed.solver_support import DESKTOP_SOLVERS
            for name in DESKTOP_SOLVERS["daetools"]:
                create_linear_solver(name)
            passed("all standard solver adapters")
            from PyQt6.QtCore import QSettings
            from PyQt6.QtWidgets import QApplication
            app = QApplication.instance() or QApplication(["MultiSolid release check"])
            from packed_bed_ui.splash import show_splash
            splash = show_splash(app)
            splash.grab().save(str(destination / 'splash.png'))
            from packed_bed_ui.branding import brand_family
            assert brand_family() == 'Maratype'
            from packed_bed_ui.window import MainWindow
            window = MainWindow(QSettings(str(destination / "settings.ini"), QSettings.Format.IniFormat))
            window.show()
            splash.finish(window)
            app.processEvents()
            window.grab().save(str(destination / "window.png"))
            window.close()
            passed("Qt window and application assets")
            passed('launch splash and bundled Maratype font')
            from packed_bed.reaction_graph import build_reaction_graph, render_svg, find_graphviz
            command = find_graphviz()
            if getattr(sys, "frozen", False):
                assert command.bundle is not None
            (destination / "graph.svg").write_bytes(render_svg(build_reaction_graph(["N2", "H2"], [], [])))
            passed("bundled Graphviz SVG rendering")
            from packed_bed_ui.inputs import resolve_documents
            documents = _documents()
            resolve_documents(documents)
            documents["run"]["simulation"]["program_mode"] = "feed_stream"
            documents["program"] = {
                "feed_stream": {"basis": "mol_per_s", "initial": {
                    "flow": 1e-8, "temperature": 300., "composition": {"N2": 1.}}, "steps": []},
                "outlet_pressure": {"initial": 100000., "steps": []}}
            resolve_documents(documents)
            passed("both program modes without example files")
            from packed_bed_ui.project import Project, read_json
            from packed_bed.reports import load_dataset
            project = Project.create(destination / "Project with spaces \u03b1", "Release check")
            for backend, name in (("daetools", "superlu"), ("compiled", "superlu"),
                                  ("compiled", "klu"), ("compiled", "band")):
                documents = _documents()
                documents["run"]["solver"].update(backend=backend, name=name)
                project.add_case(f"{backend} {name}", documents)
            for attempt in range(2):
                job = project.prepare_execution(project.cases, max_workers=2)
                with (destination / f"batch-{attempt}.log").open("wb") as log:
                    subprocess.run(_command("--project-worker", job), stdout=log, stderr=log,
                                   timeout=1200, check=True,
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
                assert read_json(job)["state"] == "completed", read_json(job)
                for case in project.cases:
                    data = load_dataset(case.run_folder / "output/results.nc")
                    assert np.allclose(data.temperature, 300., rtol=0, atol=1e-5)
                    assert np.allclose(data.gas_mole_fraction, 1., rtol=0, atol=1e-8)
                    manifest = read_json(case.run_folder / "output/manifest.json")
                    assert manifest["status"] == "success"
                    if attempt and case.documents["run"]["solver"]["backend"] == "compiled":
                        assert manifest["solver_stats"]["cache_hit"]
                passed("concurrent standard/compiled cases" if attempt == 0 else "repeat runs and project cache reuse")
            from packed_bed_ui.workbook import write_workbook
            definition = {"version": 1, "sheets": [{"name": "Temperature", "axis": "time",
                          "rows": {"mode": "all"}, "columns": [{"quantity": "outlet_temperature", "fixed": {}}]}]}
            write_workbook(project.cases[0].run_folder, definition, destination / "report.xlsx")
            from openpyxl import load_workbook
            book = load_workbook(destination / "report.xlsx")
            assert book["Temperature"].max_row >= 3
            book.close()
            passed("NetCDF read and Excel workbook export")
            from packed_bed.plotting import PLOT_REGISTRY
            from PyQt6.QtSvg import QSvgRenderer
            data = load_dataset(project.cases[0].run_folder / 'output/results.nc')
            for spec in PLOT_REGISTRY.values():
                target = destination / spec.filename
                spec.render(data, target)
                assert b'<svg' in target.read_bytes()
                assert QSvgRenderer(str(target)).isValid()
            passed('pre-made SVG plots from retained results')
            from packed_bed_ui.project_archive import export_project, import_project
            archive = destination / "project.msproject"
            export_project(project, archive)
            import_project(archive, destination / "Imported project", "Imported")
            reopened = Project.open(destination / "Imported project")
            assert len(reopened.cases) == 4
            assert not any(case.run_folder.exists() for case in reopened.cases)
            passed("project archive round trip")
            cancel_project = Project.create(destination / "Cancellation")
            documents = _documents()
            documents["run"]["solver"].update(backend="compiled", name="band")
            documents["run"]["simulation"].update(time_horizon_s=10000., reporting_interval_s=.1)
            case = cancel_project.add_case("Cancel", documents)
            job = cancel_project.prepare_execution([case])
            with (destination / "cancel.log").open("wb") as log:
                process = subprocess.Popen(_command("--project-worker", job), stdout=log, stderr=log,
                                           creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
                try:
                    deadline = time.monotonic() + 60
                    while time.monotonic() < deadline and process.poll() is None:
                        state = read_json(job)
                        if state["state"] == "running":
                            (job.parent / (".cancel-" + state["attempt_id"])).touch()
                            break
                        time.sleep(.1)
                    process.wait(timeout=60)
                    assert read_json(job)["state"] == "cancelled", read_json(job)
                finally:
                    if process.poll() is None:
                        process.kill()
                        process.wait()
            passed("worker cancellation")
        record["passed"] = True
        return 0
    except Exception:
        record.update(passed=False, error=traceback.format_exc())
        return 1
    finally:
        (destination / "result.json").write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")
